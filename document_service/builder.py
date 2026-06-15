# -*- encoding: utf-8 -*-
import io
from datetime import datetime
from typing import Dict

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from document_service.registry import get_disclaimer

BLACK = RGBColor(0, 0, 0)


def _safe_format(template: str, slots: Dict[str, str]) -> str:
    try:
        return template.format(**{k: slots.get(k, "（待填写）") for k in slots})
    except KeyError:
        return template


def _style_run(run, size: int = 12, bold: bool = False) -> None:
    run.font.name = "宋体"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = BLACK
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    r_fonts.set(qn("w:eastAsia"), "宋体")
    r_fonts.set(qn("w:ascii"), "Times New Roman")
    r_fonts.set(qn("w:hAnsi"), "Times New Roman")


def _add_paragraph(document, text: str, alignment=None, size: int = 12, bold: bool = False):
    p = document.add_paragraph()
    if alignment is not None:
        p.alignment = alignment
    run = p.add_run(text)
    _style_run(run, size=size, bold=bold)
    return p


def _apply_black_to_document(document: Document) -> None:
    for paragraph in document.paragraphs:
        for run in paragraph.runs:
            size = run.font.size.pt if run.font.size is not None else 12
            _style_run(run, size=int(size), bold=run.font.bold)


def build_document_bytes(
    doc_type: Dict,
    slots: Dict[str, str],
) -> bytes:
    """Build application letter only; reasoning and regulation excerpts stay in UI."""
    document = Document()
    style = document.styles["Normal"]
    style.font.name = "宋体"
    style.font.size = Pt(12)
    style.font.color.rgb = BLACK

    title = doc_type.get("document_title", "办事申请说明")
    heading = document.add_heading(title, level=0)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in heading.runs:
        _style_run(run, size=16, bold=True)

    _add_paragraph(
        document,
        "（系统生成草稿，非学校官方空白表格）",
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        size=10,
    )

    addressee = doc_type.get("addressee", "学院/教务处")
    _add_paragraph(document, f"{addressee}：")

    _add_paragraph(
        document,
        f"姓名：{slots.get('student_name', '')}    学号：{slots.get('student_id', '')}",
    )
    if slots.get("college"):
        line = f"学院：{slots.get('college', '')}"
        if slots.get("major"):
            line += f"    专业：{slots.get('major', '')}"
        _add_paragraph(document, line)

    body_template = doc_type.get("body_template", "")
    if body_template:
        body = _safe_format(body_template, slots)
    else:
        body = (
            f"本人因{slots.get('reason', '相关事由')}，现提出"
            f"{title}，请予审核。"
        )
    _add_paragraph(document, body)

    document.add_paragraph("")

    attachments = list(doc_type.get("default_attachments") or [])
    if slots.get("attachments_note"):
        attachments.append(slots["attachments_note"])
    if attachments:
        _add_paragraph(document, "随附材料：", bold=True)
        for item in attachments:
            p = document.add_paragraph(style="List Bullet")
            run = p.add_run(item)
            _style_run(run)

    document.add_paragraph("")
    contact_line = f"联系方式：{slots.get('contact', '')}"
    if slots.get("phone"):
        contact_line += f"（{slots.get('phone', '')}）"
    _add_paragraph(document, contact_line)
    _add_paragraph(document, f"申请日期：{datetime.now().strftime('%Y年%m月%d日')}")
    _add_paragraph(document, "申请人（签字）：__________")

    document.add_paragraph("")
    disclaimer = get_disclaimer()
    if disclaimer:
        _add_paragraph(document, disclaimer, size=9)

    _apply_black_to_document(document)

    buffer = io.BytesIO()
    document.save(buffer)
    buffer.seek(0)
    return buffer.read()


def suggest_filename(doc_type: Dict, student_name: str) -> str:
    title = doc_type.get("document_title", "办事文书")
    name = student_name.strip() or "未命名"
    date = datetime.now().strftime("%Y%m%d")
    safe_title = title.replace("/", "_")
    return f"{safe_title}_{name}_{date}.docx"
